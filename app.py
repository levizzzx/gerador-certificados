import os
import re
import io
import docx
import zipfile
import pandas as pd
import streamlit as st
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

# configuração da página no streamlit
st.set_page_config(
    page_title="Gerador de Certificados",
    layout="centered"
)

def extrair_alunos_word(file_bytes):
    # lê arquivos .docx e extrai nome e cpf usando expressões regulares
    doc = docx.Document(io.BytesIO(file_bytes))
    texto = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    
    padrao = re.compile(
        r"Nome\s*completo:\s*(?P<nome>[^\n]+)\n(?:Função:[^\n]+\n)?CPF:\s*(?P<cpf>[^\n]+)",
        re.IGNORECASE
    )
    
    alunos = []
    for match in padrao.finditer(texto):
        alunos.append({
            "nome": match.group("nome").strip(),
            "cpf": match.group("cpf").strip()
        })
    return alunos

def extrair_alunos_excel(file_bytes, is_csv=False):
    # lê planilhas .xlsx ou .csv e extrai as colunas nome e cpf
    if is_csv:
        df = pd.read_csv(io.BytesIO(file_bytes))
    else:
        df = pd.read_excel(io.BytesIO(file_bytes))
    
    alunos = []
    for _, row in df.iterrows():
        nome_col = next((col for col in df.columns if 'nome' in col.lower()), None)
        cpf_col = next((col for col in df.columns if 'cpf' in col.lower()), None)
        
        nome = str(row[nome_col]).strip() if nome_col else str(row.get('Nome', '')).strip()
        cpf = str(row[cpf_col]).strip() if cpf_col else str(row.get('CPF', '')).strip()
        
        if nome and nome.lower() != 'nan':
            alunos.append({"nome": nome, "cpf": cpf})
    return alunos


def processar_frame_de_texto(shape, text_frame, nome, cpf):
    # substitui marcações e ajusta exclusivamente a largura da caixa do rodapé
    
    eh_caixa_rodape = False

    # identifica se a caixa está na metade inferior do slide (rodapé)
    posicao_top = getattr(shape, 'top', 0)
    if posicao_top > Inches(3.5):
        for paragraph in text_frame.paragraphs:
            texto_p = paragraph.text
            if "Portador" in texto_p or ("CPF" in texto_p and "CONCLUIU" not in texto_p) or re.search(r'X{3,}', texto_p):
                eh_caixa_rodape = True
                break

    for paragraph in text_frame.paragraphs:
        texto_p = paragraph.text
        if not texto_p.strip():
            continue

        # 1. substituição do cpf
        padroes_cpf = ["XXX.XXX.XXX-XX", "{CPF}", "[CPF]", "<CPF>", "{cpf}", "[cpf]"]
        for p_cpf in padroes_cpf:
            if p_cpf in texto_p:
                for run in paragraph.runs:
                    if p_cpf in run.text:
                        run.text = run.text.replace(p_cpf, cpf)
                if p_cpf in paragraph.text and not any(p_cpf in r.text for r in paragraph.runs):
                    if paragraph.runs:
                        paragraph.runs[0].text = paragraph.text.replace(p_cpf, cpf)
                        for r in paragraph.runs[1:]:
                            r.text = ""

        # 2. substituição das sequências de x ou tags de nome
        padroes_nome_tags = ["{NOME}", "[NOME]", "<NOME>", "{nome}", "[nome]"]
        for tag in padroes_nome_tags:
            if tag in texto_p:
                for run in paragraph.runs:
                    if tag in run.text:
                        run.text = run.text.replace(tag, nome)

        # se houver sequências de x (nome)
        if re.search(r'X{3,}', texto_p, re.IGNORECASE):
            for run in paragraph.runs:
                if re.search(r'X{3,}', run.text, re.IGNORECASE):
                    run.text = re.sub(r'X{3,}', nome, run.text, flags=re.IGNORECASE)
            
            if re.search(r'X{3,}', paragraph.text, re.IGNORECASE) and not any(re.search(r'X{3,}', r.text, re.IGNORECASE) for r in paragraph.runs):
                if paragraph.runs:
                    paragraph.runs[0].text = re.sub(r'X{3,}', nome, paragraph.text, flags=re.IGNORECASE)
                    for r in paragraph.runs[1:]:
                        r.text = ""

    # alarga apenas se for a caixa do rodapé
    if eh_caixa_rodape:
        try:
            text_frame.word_wrap = False
            text_frame.margin_left = Inches(0)
            text_frame.margin_right = Inches(0)
            
            # expande a largura da caixa do rodapé em 1.2 polegadas
            shape.width = shape.width + Inches(1.2)
            
            # reduz levemente a fonte se o nome for grande no rodapé
            if len(nome) > 25:
                for p in text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size > Pt(8.5):
                            r.font.size = Pt(8.5)
        except Exception:
            pass


def processar_forma_recursiva(shape, nome, cpf):
    # navega por formas sem alterar caixas principais do meio ou grupos
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub_shape in shape.shapes:
            processar_forma_recursiva(sub_shape, nome, cpf)
            
    elif shape.has_text_frame:
        processar_frame_de_texto(shape, shape.text_frame, nome, cpf)

    elif shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                if cell.text_frame:
                    processar_frame_de_texto(shape, cell.text_frame, nome, cpf)


# --- interface web (streamlit) ---

st.title("Gerador de Certificados")
st.caption("Envie a lista de participantes e o modelo em PowerPoint (.pptx) para processamento automático.")

st.markdown("---")

col1, col2 = st.columns(2)

with col1:
    st.subheader("1. Lista de Alunos")
    file_alunos = st.file_uploader(
        "Arquivo de dados (.docx, .xlsx, .csv)", 
        type=["docx", "xlsx", "csv"], 
        key="alunos_file"
    )

with col2:
    st.subheader("2. Modelo")
    file_modelo = st.file_uploader(
        "Modelo do certificado (.pptx)", 
        type=["pptx"], 
        key="modelo_file"
    )

if file_alunos is not None and file_modelo is not None:
    bytes_alunos = file_alunos.read()
    nome_arq_alunos = file_alunos.name.lower()
    bytes_modelo = file_modelo.read()
    
    alunos = []
    if nome_arq_alunos.endswith('.docx'):
        alunos = extrair_alunos_word(bytes_alunos)
    elif nome_arq_alunos.endswith('.csv'):
        alunos = extrair_alunos_excel(bytes_alunos, is_csv=True)
    elif nome_arq_alunos.endswith(('.xlsx', '.xls')):
        alunos = extrair_alunos_excel(bytes_alunos, is_csv=False)

    if not alunos:
        st.error("Nenhum participante identificado no arquivo enviado.")
    else:
        st.info(f"{len(alunos)} participante(s) localizado(s).")
        
        df_preview = pd.DataFrame(alunos)
        st.dataframe(df_preview, use_container_width=True)

        if st.button("Gerar Certificados", type="primary"):
            zip_buffer = io.BytesIO()
            
            with zipfile.ZipFile(zip_buffer, "w") as zip_file:
                for aluno in alunos:
                    nome_aluno = aluno["nome"]
                    cpf_aluno = aluno["cpf"]

                    prs = Presentation(io.BytesIO(bytes_modelo))

                    for slide in prs.slides:
                        for shape in slide.shapes:
                            processar_forma_recursiva(shape, nome_aluno, cpf_aluno)

                    cert_buffer = io.BytesIO()
                    prs.save(cert_buffer)
                    cert_buffer.seek(0)
                    
                    nome_cert = f"Certificado_{nome_aluno.replace(' ', '_')}.pptx"
                    zip_file.writestr(nome_cert, cert_buffer.getvalue())

            zip_buffer.seek(0)
            
            st.download_button(
                label="Download dos Certificados (.ZIP)",
                data=zip_buffer,
                file_name="Certificados_Gerados.zip",
                mime="application/zip"
            )